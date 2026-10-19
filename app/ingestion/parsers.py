
import io
from datetime import datetime
from pathlib import Path

import pymupdf
import pytesseract
from PIL import Image
from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.schemas import ParsedDocument


def _get_last_updated(path: Path) -> str:
    """Return the file's last-modified timestamp."""
    return datetime.fromtimestamp(
        path.stat().st_mtime
    ).isoformat(timespec="seconds")


def _make_record(
    path: Path,
    document_id: str,
    file_type: str,
    text: str,
    last_updated: str,
    page: int | None = None,
    section: str | None = None,
    extraction_method: str = "native_text",
) -> ParsedDocument:
    """Create a consistently structured document record."""
    return ParsedDocument(
        document_id=document_id,
        source=path.name,
        file_type=file_type,
        text=text.strip(),
        page=page,
        section=section,
        last_updated=last_updated,
        metadata={"extraction_method": extraction_method},
    )


def _escape_table_cell(value: str) -> str:
    """Escape pipe characters so cells remain valid Markdown."""
    return value.strip().replace("|", r"\|").replace("\n", " ")


def _rows_to_markdown(rows: list[list[str]]) -> str:
    """Convert rows into a Markdown table with a header separator."""
    cleaned_rows = [
        [_escape_table_cell(cell) for cell in row]
        for row in rows
    ]

    cleaned_rows = [
        row for row in cleaned_rows if any(cell for cell in row)
    ]

    if not cleaned_rows:
        return ""

    column_count = max(len(row) for row in cleaned_rows)

    normalized_rows = [
        row + [""] * (column_count - len(row))
        for row in cleaned_rows
    ]

    header = normalized_rows[0]
    separator = ["---"] * column_count

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]

    for row in normalized_rows[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def _html_table_to_markdown(table_element) -> str:
    """Convert an HTML table into Markdown."""
    rows: list[list[str]] = []

    for tr in table_element.find_all("tr"):
        cells = tr.find_all(["th", "td"], recursive=False)

        if cells:
            rows.append(
                [
                    cell.get_text(" ", strip=True)
                    for cell in cells
                ]
            )

    return _rows_to_markdown(rows)


def _extract_pdf_tables(page) -> list[str]:
    """
    Extract detectable PDF tables using PyMuPDF.

    Returns an empty list if table detection is unavailable or fails.
    """
    try:
        finder = page.find_tables()
        tables: list[str] = []

        for table in finder.tables:
            extracted = table.extract()

            if not extracted:
                continue

            markdown_table = _rows_to_markdown(
                [
                    [str(cell or "") for cell in row]
                    for row in extracted
                ]
            )

            if markdown_table:
                tables.append(markdown_table)

        return tables

    except Exception as exc:
        print(f"PDF table extraction warning: {exc}")
        return []


def _iter_docx_blocks(document):
    """Yield DOCX paragraphs and tables in document order."""
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _extract_docx_content(document) -> str:
    """Extract DOCX paragraphs, headings, and tables in order."""
    parts: list[str] = []

    for block in _iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            content = block.text.strip()

            if not content:
                continue

            style_name = (
                block.style.name
                if block.style is not None
                else ""
            )

            if style_name.startswith("Heading"):
                match = __import__("re").search(
                    r"(\d+)", style_name
                )
                level = int(match.group(1)) if match else 1
                level = max(1, min(level, 6))
                parts.append(f"{'#' * level} {content}")
            else:
                parts.append(content)

        elif isinstance(block, Table):
            rows = []

            for row in block.rows:
                rows.append(
                    [
                        cell.text.strip().replace("\n", " ")
                        for cell in row.cells
                    ]
                )

            markdown_table = _rows_to_markdown(rows)

            if markdown_table:
                parts.append(markdown_table)

    return "\n\n".join(parts).strip()


def parse_document(file_path: str | Path) -> list[ParsedDocument]:
    """
    Parse PDF, DOCX, Markdown, or HTML documents.

    Features:
    - PDF native text extraction with OCR fallback.
    - PDF table detection when supported by PyMuPDF.
    - DOCX paragraph, heading, and table extraction.
    - Markdown heading and table preservation.
    - HTML table conversion and removal of scripts/styles.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Document not found: {path}")

    if not path.is_file():
        raise ValueError(f"Expected a file, received: {path}")

    suffix = path.suffix.lower()
    document_id = path.stem
    last_updated = _get_last_updated(path)
    records: list[ParsedDocument] = []

    # 1. PDF: extract native text, tables, and OCR if necessary.
    if suffix == ".pdf":
        with pymupdf.open(path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()
                extraction_method = "native_text"

                if not text:
                    print(f"Page {page_number}: Running OCR...")

                    pixmap = page.get_pixmap(
                        matrix=pymupdf.Matrix(3, 3),
                        alpha=False,
                    )

                    image = Image.open(
                        io.BytesIO(pixmap.tobytes("png"))
                    )

                    text = pytesseract.image_to_string(
                        image,
                        lang="eng",
                    ).strip()

                    extraction_method = "ocr"

                    print(
                        f"Page {page_number}: OCR extracted "
                        f"{len(text)} characters."
                    )

                # Only use native PDF table detection when native
                # text is available. OCR-only tables need a separate
                # layout-aware OCR strategy.
                if extraction_method == "native_text":
                    tables = _extract_pdf_tables(page)

                    if tables:
                        text = (
                            text + "\n\n"
                            + "\n\n".join(tables)
                        ).strip()
                        extraction_method = "native_text_with_tables"

                if text:
                    records.append(
                        _make_record(
                            path=path,
                            document_id=document_id,
                            file_type="pdf",
                            text=text,
                            last_updated=last_updated,
                            page=page_number,
                            extraction_method=extraction_method,
                        )
                    )
                else:
                    print(
                        f"Page {page_number}: No text found "
                        "using native extraction or OCR."
                    )

    # 2. DOCX: preserve heading levels and tables in document order.
    elif suffix == ".docx":
        doc = Document(path)
        text = _extract_docx_content(doc)

        if text:
            records.append(
                _make_record(
                    path=path,
                    document_id=document_id,
                    file_type="docx",
                    text=text,
                    last_updated=last_updated,
                    extraction_method="docx_with_tables",
                )
            )

    # 3. Markdown: keep the original Markdown syntax.
    elif suffix in {".md", ".markdown"}:
        text = path.read_text(encoding="utf-8-sig").strip()

        if text:
            records.append(
                _make_record(
                    path=path,
                    document_id=document_id,
                    file_type="md",
                    text=text,
                    last_updated=last_updated,
                    extraction_method="markdown",
                )
            )

    # 4. HTML: preserve headings and convert tables to Markdown.
    elif suffix in {".html", ".htm"}:
        raw_html = path.read_text(
            encoding="utf-8",
            errors="replace",
        )

        soup = BeautifulSoup(raw_html, "html.parser")

        for unwanted in soup(
            ["script", "style", "noscript"]
        ):
            unwanted.decompose()

        parts: list[str] = []

        for element in soup.find_all(
            ["h1", "h2", "h3", "h4", "h5", "h6",
             "p", "table", "ul", "ol", "pre", "blockquote"]
        ):
            if element.name == "table":
                content = _html_table_to_markdown(element)
            elif element.name.startswith("h"):
                level = int(element.name[1])
                heading = element.get_text(" ", strip=True)
                content = f"{'#' * level} {heading}"
            else:
                content = element.get_text(" ", strip=True)

            if content.strip():
                parts.append(content.strip())

        text = "\n\n".join(parts).strip()

        if text:
            records.append(
                _make_record(
                    path=path,
                    document_id=document_id,
                    file_type="html",
                    text=text,
                    last_updated=last_updated,
                    extraction_method="html_with_tables",
                )
            )

    else:
        raise ValueError(
            f"Unsupported file type: {suffix}. "
            "Supported types: PDF, DOCX, Markdown, HTML."
        )

    print(
        f"Finished parsing {path.name}: "
        f"{len(records)} record(s) created."
    )

    return records
